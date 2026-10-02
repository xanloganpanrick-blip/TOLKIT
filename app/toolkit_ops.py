"""Shared server operations for the original Toolkit sections."""
from pathlib import Path
from datetime import datetime, timedelta
import csv, heapq, io, math, random, re, zipfile
import polars as pl
from . import table_ops as t

OPERATIONS = {
    "convert_csv","clean","normalize","remove7","dedupe","filter","sort","merge","split","stats",
    "phones","geo","regions","max_filter","row_names","fssp","distribute","take","addresses","group_split",
}

# Phone candidates may contain spaces, parentheses or hyphens. Never join different numbers.
PHONE_RE = re.compile(r"(?<!\d)(?:\+?[78][\s().-]*)?(?:\d[\s().-]*){9}\d(?!\d)")

def extract_phones(value):
    text = str(value or "").strip()
    result=[]
    # Adjacent unformatted numbers also occur without any separators.
    for token in re.split(r"[;,|\n\r/]+", text):
        compact=re.sub(r"\s+", "", token)
        if compact.isdigit() and len(compact)>15:
            if len(compact)%11 == 0 and all(compact[i] in "78" for i in range(0,len(compact),11)):
                candidates=[compact[i:i+11] for i in range(0,len(compact),11)]
            elif len(compact)%10 == 0:
                candidates=[compact[i:i+10] for i in range(0,len(compact),10)]
            else: candidates=[]
        else:
            candidates=[m.group() for m in PHONE_RE.finditer(token)]
        for candidate in candidates:
            digits=re.sub(r"\D","",candidate)
            if len(digits)==10: digits="7"+digits
            elif len(digits)==11 and digits[0]=="8": digits="7"+digits[1:]
            if len(digits)==11 and digits[0]=="7" and digits not in result:
                result.append(digits)
    return result

def phone_key(value):
    numbers=extract_phones(value)
    return numbers[0][-10:] if numbers else ""

def phone_column(df, options=None):
    requested=(options or {}).get("phone_column","").strip()
    if requested:
        if requested not in df.columns: raise ValueError("Колонка номера не найдена: "+requested)
        return requested
    known=t._find_col(df.columns,t.PHONE_ALIASES)
    if known: return known
    # Infer from a sample without scanning unrelated address/date cells in every row.
    best,score=None,0
    for col in df.columns:
        if any(word in col.lower() for word in ("адрес","address","снилс","инн","date","дата")): continue
        values=df[col].head(200).to_list()
        valid=sum(bool(extract_phones(value)) for value in values)
        if valid>score: best,score=col,valid
    if best and score/max(1,min(200,df.height))>=.4: return best
    raise ValueError("Не найдена колонка с номерами. Укажи её в настройках.")

ADDRESS_LABELS=[
    ("city",r"город|г\.?","г."),
    ("street",r"улица|ул\.?","ул."),
    ("house",r"дом|д\.?","д."),
    ("corpus",r"корпус|корп\.?|к\.?","к."),
    ("building",r"строение|строен\.?|стр\.?|с\.?","с."),
    ("letter",r"литера|лит\.?","лит."),
    ("flat",r"квартира|кв\.?","кв."),
    ("street",r"проспект|просп\.?|пр-т","пр-т"),
    ("street",r"переулок|пер\.?","пер."),
    ("street",r"шоссе|ш\.?","ш."),
    ("street",r"набережная|наб\.?","наб."),
]
ADDRESS_RE=re.compile(r"(?<!\w)("+ "|".join(pattern for _,pattern,_ in ADDRESS_LABELS)+r")(?=\s|\d|$)\s*",re.I)

def format_address(value):
    text=re.sub(r"\s+"," ",str(value or "")).strip(" ,;")
    if not text: return "",False
    text=re.sub(r"(?<!\w)(г|ул|д|к|с|кв|стр|корп|лит|пер|просп|наб)\.(?=\S)",r"\1. ",text,flags=re.I)
    # Postfix house/corpus markers are common in database exports.
    text=re.sub(r"(\d+[А-Яа-яA-Za-z]?)\s+дом\b(?=\s*[,;]|$)",r"д. \1",text,flags=re.I)
    text=re.sub(r"(\d+)\s+корпус\b(?=\s*[,;]|$)",r"к. \1",text,flags=re.I)
    text=re.sub(r"(?<=\d)([кс])(?=\d)",r" \1. ",text,flags=re.I)
    found=list(ADDRESS_RE.finditer(text)); parts={}; labels={}; extras=[]; uncertain=False
    for i,match in enumerate(found):
        label=match.group(1)
        key,prefix=next((key,prefix) for key,pattern,prefix in ADDRESS_LABELS
                        if re.fullmatch(pattern,label,re.I))
        item=text[match.end():found[i+1].start() if i+1<len(found) else len(text)].strip(" ,;")
        if item:
            if key in parts: extras.append(prefix+" "+item); uncertain=True
            else: parts[key]=item; labels[key]=prefix
    before=text[:found[0].start()].strip(" ,;") if found else text
    if before:
        chunks=[x.strip() for x in before.split(",") if x.strip()]
        # Postal code / region are retained rather than silently discarded.
        candidates=[x for x in chunks if not re.search(r"\d|обл\.|область|край|респ",x,re.I)]
        if "city" not in parts and candidates:
            parts["city"]=candidates[0]; labels["city"]="г."
            chunks.remove(candidates[0])
        if not found and chunks and "street" not in parts:
            parts["street"]=chunks.pop(0); labels["street"]="ул."
            if chunks and re.fullmatch(r"\d+[А-Яа-яA-Za-z]?(?:/\d+)?",chunks[0]):
                parts["house"]=chunks.pop(0); labels["house"]="д."
        extras=chunks+extras
        uncertain=uncertain or bool(chunks)
    if "street" in parts and "house" not in parts:
        tail=re.match(r"^(.+?)\s*,\s*(\d+[А-Яа-яA-Za-z]?(?:/\d+)?)$",parts["street"])
        if tail: parts["street"],parts["house"]=tail.groups(); labels["house"]="д."
    def title(item):
        return item.title() if item.isupper() or item.islower() else item
    city=title(parts.get("city","")); street=title(parts.get("street",""))
    result=[]
    if city: result.append("г. "+city)
    if street: result.append(labels.get("street","ул.")+" "+street)
    for key,prefix in (("house","д."),("corpus","к."),("building","с."),("letter","лит."),("flat","кв.")):
        if parts.get(key): result.append(prefix+" "+parts[key])
    if extras: result.extend(extras)
    return ", ".join(result) if result else text, uncertain or not(city and street and parts.get("house"))

def addresses(df, options):
    col=options.get("address_column") or t._find_col(df.columns,{"адрес","адресс","address"})
    if not col or col not in df.columns: raise ValueError("Не найдена колонка адреса")
    parsed=[format_address(value) for value in df[col].to_list()]
    return df.with_columns(pl.Series(col,[x[0] for x in parsed]),
                           pl.Series("Адрес требует проверки",[x[1] for x in parsed]))

def date_dmy(value):
    text=str(value or "").strip()
    for fmt in ("%d.%m.%Y","%d/%m/%Y","%Y-%m-%d","%Y-%m-%d %H:%M:%S","%d.%m.%y"):
        try: return datetime.strptime(text,fmt).strftime("%d.%m.%Y")
        except ValueError: pass
    # Excel serial dates are converted only in a recognised date column.
    try:
        serial=float(text)
        if 10000<=serial<=80000: return (datetime(1899,12,30)+timedelta(days=serial)).strftime("%d.%m.%Y")
    except ValueError: pass
    return text

def geo(df, options):
    col=phone_column(df,options)
    extra=[name for name in df.columns if name!=col and
           (t._find_col([name],t.PHONE_ALIASES) or name.startswith("col_")) and
           any(extract_phones(value) for value in df[name].head(100).to_list())]
    if extra:
        def combined(row):
            result=[]
            for name in [col,*extra]:
                for phone in extract_phones(row[name]):
                    if phone not in result:result.append(phone)
            return result
        df=df.with_columns(pl.struct([col,*extra]).map_elements(combined,return_dtype=pl.List(pl.String)).alias(col))
    else:
        df=df.with_columns(pl.col(col).map_elements(extract_phones,return_dtype=pl.List(pl.String)).alias(col))
    df=df.explode(col)
    fio=t._find_col(df.columns,t.FIO_ALIASES)
    date=t._find_col(df.columns,t.DATE_ALIASES)
    if fio:
        df=df.with_columns(t.normalize_fio_expr(fio))
        if options.get("exclude_names"):
            df=df.filter(~pl.col(fio).str.to_lowercase().str.contains("оглы|кызы").fill_null(False))
        df=df.with_columns((pl.col(fio).str.count_matches(r"\S+")<3).alias("__incomplete"))
        df=df.sort("__incomplete",maintain_order=True).drop("__incomplete")
    if date: df=df.with_columns(pl.col(date).map_elements(date_dmy,return_dtype=pl.String).alias(date))
    if options.get("format_addresses") and t._find_col(df.columns,{"адрес","адресс","address"}):
        df=addresses(df,options)
    return df

def write_table(df,path,options):
    path=path.parent/(path.name+(".xlsx" if options.get("output_format")=="xlsx" else ".csv"))
    path.parent.mkdir(parents=True,exist_ok=True)
    if path.suffix==".xlsx":
        if df.height>1_048_575 or df.width>16_384: raise ValueError("Таблица превышает лимит Excel. Выбери CSV.")
        from openpyxl import Workbook
        wb=Workbook(write_only=True); ws=wb.create_sheet("Данные"); ws.append(df.columns)
        from openpyxl.cell import WriteOnlyCell
        for row in df.iter_rows():
            cells=[]
            for value in row:
                cell=WriteOnlyCell(ws,value=value)
                if isinstance(value,str): cell.data_type="s"  # preserve literal values, not executable formulas
                cells.append(cell)
            ws.append(cells)
        wb.save(path); wb.close()
    else: df.write_csv(path,include_bom=True,separator=options.get("csv_separator",","))
    return path

def archive(paths,target):
    with zipfile.ZipFile(target,"w",compression=zipfile.ZIP_DEFLATED,allowZip64=True) as z:
        for path in paths: z.write(path,path.name)
    return target

def reference_numbers(paths):
    values={}
    for path in paths:
        if path.suffix.lower()==".txt":
            for line in t.decode_text(path).splitlines():
                matches=list(PHONE_RE.finditer(line))
                if not matches: continue
                # Separate phone from MAX marker. The marker never affects the phone key.
                first=matches[0]; key=phone_key(first.group())
                if key: values.setdefault(key,line[first.end():].strip(" ,;\t"))
        else:
            frame=t.read_any(path); col=phone_column(frame)
            max_col=t._find_col(frame.columns,{"max","макс"})
            for row in frame.select([col]+([max_col] if max_col and max_col!=col else [])).iter_rows():
                for phone in extract_phones(row[0]):
                    values.setdefault(phone[-10:],str(row[1] or "") if len(row)>1 else "")
    if not values: raise ValueError("В справочнике нет корректных номеров")
    return pl.DataFrame({"__phone":list(values),"__MAX":list(values.values())})

def max_filter(df,refs,options):
    if (options.get("matches_only") or options.get("download_unmatched")) and not refs:
        raise ValueError("Добавь файл номеров для сверки")
    df=geo(df,{**options,"exclude_names":False})
    col=phone_column(df,options)
    # GEO already returned canonical phones; keep matching on the native Polars path.
    df=df.with_columns(pl.col(col).str.slice(-10).alias("__phone"))
    if refs:
        lookup=reference_numbers(refs)
        df=df.join(lookup.with_columns(pl.lit(True).alias("__matched")),on="__phone",how="left",maintain_order="left")
        df=df.with_columns(pl.col("__MAX").fill_null("").alias("MAX"))
        matched=df.filter(pl.col("__matched").fill_null(False))
        unmatched=df.filter(~pl.col("__matched").fill_null(False))
    else: matched=df; unmatched=df.head(0)
    def finish(frame):
        frame=frame.drop([name for name in ("__phone","__MAX","__matched") if name in frame.columns])
        if options.get("remove7"):
            frame=frame.with_columns(pl.col(col).str.strip_prefix("7").alias(col))
        if options.get("replace_century"):
            date=t._find_col(frame.columns,t.DATE_ALIASES)
            if date: frame=frame.with_columns(pl.col(date).str.replace(r"\.20(\d{2})$",r".19$1").alias(date))
        if options.get("dedupe"): frame=frame.unique(subset=[col],keep="first",maintain_order=True)
        return frame
    return finish(matched if options.get("matches_only") else df),finish(unmatched)

def spread_indices(length,need):
    if not length or not need: return []
    first=random.randrange(length); heap=[]; selected=[first]
    def push(left,right):
        if right-left<=1:return
        candidate=0 if left<0 else length-1 if right>=length else (left+right)//2
        distance=right-candidate if left<0 else candidate-left if right>=length else min(candidate-left,right-candidate)
        heapq.heappush(heap,(-distance,candidate,left,right))
    push(-1,first); push(first,length)
    while heap and len(selected)<min(length,need):
        _,candidate,left,right=heapq.heappop(heap); selected.append(candidate)
        push(left,candidate); push(candidate,right)
    return selected

# The original prefix mapping is deliberately not used: mobile codes span multiple regions.
def regions(df,refs,options):
    if not refs: raise ValueError("Добавь CSV-справочник диапазонов: код, от, до, регион")
    tables=[t.read_any(path) for path in refs]
    reference=pl.concat(tables,how="diagonal_relaxed")
    def field(aliases):
        name=t._find_col(reference.columns,aliases)
        if not name: raise ValueError("В справочнике не найдено поле: "+"/".join(aliases))
        return name
    code=field({"авс/ def","авс/def","abc/def","код","code"})
    start=field({"от","from","начало"}); end=field({"до","to","конец"})
    region=field({"регион","region"})
    ranges={}
    for row in reference.select(code,start,end,region).iter_rows():
        try: ranges.setdefault(str(row[0]).strip(),[]).append((int(row[1]),int(row[2]),str(row[3])))
        except (ValueError,TypeError): continue
    if not ranges: raise ValueError("Справочник не содержит диапазонов")
    for items in ranges.values():items.sort()
    from bisect import bisect_right
    starts={key:[x[0] for x in values] for key,values in ranges.items()}
    def locate(value):
        phone=phone_key(value)
        if not phone:return "Не определён"
        key,number=phone[:3],int(phone[3:])
        items=ranges.get(key,[]); index=bisect_right(starts.get(key,[]),number)-1
        if index>=0 and items[index][0]<=number<=items[index][1]:return items[index][2]
        return "Не определён"
    col=phone_column(df,options)
    return df.with_columns(pl.col(col).map_elements(locate,return_dtype=pl.String).alias("Регион РФ"))

def run_operation(operation,paths,refs,options,work,progress):
    if operation not in OPERATIONS: raise ValueError("Неизвестная операция")
    if options.get("output_format","csv") not in ("csv","xlsx"): raise ValueError("Выбери CSV или XLSX")
    if options.get("csv_separator",",") not in (",",";","\t","|"): raise ValueError("Неверный разделитель CSV")
    outputs=[]
    def save(df,name): outputs.append(write_table(df,work/name,options))
    def safe(value): return re.sub(r'[<>:"/\\|?*\x00-\x1f]',"_",str(value))[:120] or "result"
    if operation=="phones":
        numbers=set()
        for i,path in enumerate(paths):
            frame=t.read_any(path); col=phone_column(frame,options)
            for value in frame[col].to_list():numbers.update(extract_phones(value))
            progress(20+int((i+1)/len(paths)*70))
        out=work/"numbers.txt"; out.write_text("\n".join(sorted(numbers)),encoding="utf-8"); return [out]
    if operation=="merge":
        frames=[]; canonical={}
        for i,path in enumerate(paths,1):
            for _,frame in t.read_tables(path):
                rename={}
                for col in frame.columns:
                    key=t._norm_header(col); canonical.setdefault(key,col); rename[col]=canonical[key]
                frame=frame.rename(rename).with_columns(pl.lit(str(i)).alias("№ таблицы"))
                frames.append(frame)
        save(pl.concat(frames,how="diagonal_relaxed"),"merged")
    elif operation=="take":
        frames=[t.read_tables(path,first_sheet=True)[0][1] for path in paths]
        target=int(options.get("count",1000))
        if target<1:raise ValueError("Количество строк должно быть больше нуля")
        allocations=[0]*len(frames); remaining=min(target,sum(frame.height for frame in frames))
        # Fill quotas in rounds in O(files * files), rather than visiting every row.
        active=[i for i,frame in enumerate(frames) if frame.height]
        while remaining and active:
            share=max(1,remaining//len(active))
            for i in active:
                n=min(share,frames[i].height-allocations[i],remaining); allocations[i]+=n; remaining-=n
            active=[i for i in active if allocations[i]<frames[i].height]
        for i,(path,frame) in enumerate(zip(paths,frames)):
            indices=spread_indices(frame.height,allocations[i])
            # Use an index membership mask, keeping original order in the remainder.
            indexed=frame.with_row_index("__index")
            selected=indexed.filter(pl.col("__index").is_in(indices)).drop("__index")
            rest=indexed.filter(~pl.col("__index").is_in(indices)).drop("__index")
            save(selected,f"{i+1}_кусок_{path.stem}"); save(rest,f"{i+1}_остаток_{path.stem}")
    else:
        for i,path in enumerate(paths,1):
            name=safe(path.stem)
            if sum(other.stem.casefold()==path.stem.casefold() for other in paths)>1:
                name=str(i)+"_"+name
            if operation=="fssp" and options.get("output_format")=="xlsx" and path.suffix.lower() in {".xlsx",".xlsm"}:
                from openpyxl import load_workbook
                wb=load_workbook(path)
                changed=0
                try:
                    for ws in wb.worksheets:
                        for row in ws.iter_rows(min_row=1,max_row=min(20,ws.max_row)):
                            for cell in row:
                                if t._norm_header(cell.value)=="номер":
                                    cell.value="Номер телефона"; changed+=1
                    if not changed: raise ValueError("Заголовок «Номер» не найден в первых 20 строках")
                    output=work/(name+"_ФССП.xlsx");wb.save(output);outputs.append(output)
                finally: wb.close()
                progress(20+int(i/len(paths)*75));continue
            frame=t.read_tables(path,first_sheet=True)[0][1] if operation in {"distribute","row_names"} else t.read_any(path,search_header=operation in {"fssp","remove7"})
            if operation=="convert_csv": outputs.append(write_table(frame,work/name,{**options,"output_format":"csv"}))
            elif operation=="clean":save(t.clean(frame),name+"_clean")
            elif operation=="normalize":save(t.normalize(frame),name+"_normalized")
            elif operation=="remove7":
                col=phone_column(frame,options)
                save(frame.with_columns(pl.col(col).cast(pl.String).str.strip_chars_start().str.strip_prefix("7").alias(col)),name+"_minus7")
            elif operation=="dedupe":
                keys=options.get("keys") or None
                if keys and any(key not in frame.columns for key in keys): raise ValueError("Колонка для удаления дублей не найдена")
                save(frame.unique(subset=keys,keep="first",maintain_order=True),name+"_dedup")
            elif operation=="filter":save(t.filter_rows(frame,options.get("column",""),options.get("value",""),options.get("contains",False)),name+"_filtered")
            elif operation=="sort":save(t.sort_rows(frame,options.get("column",""),options.get("descending",False)),name+"_sorted")
            elif operation=="geo":save(geo(frame,options),name+"_ГЕО")
            elif operation=="addresses":save(addresses(frame,options),name+"_адреса")
            elif operation=="max_filter":
                result,unmatched=max_filter(frame,refs,options);save(result,name+"_MAX")
                if options.get("download_unmatched"):save(unmatched,name+"_несовпадения")
            elif operation=="fssp":
                col=t._find_col(frame.columns,{"номер"})
                if not col:raise ValueError("Не найден заголовок «Номер»")
                if "Номер телефона" in frame.columns and col!="Номер телефона":raise ValueError("Колонка «Номер телефона» уже существует")
                save(frame.rename({col:"Номер телефона"}),name+"_ФССП")
            elif operation=="row_names":
                renamed=f"{frame.height} "+re.sub(r"^\d+\s+","",path.name)
                import shutil
                out=work/renamed;shutil.copy2(path,out);outputs.append(out)
            elif operation in {"split","distribute","group_split","regions"}:
                if operation=="split":
                    limit=int(options.get("max_rows",500000))
                    if limit<1:raise ValueError("Размер части должен быть больше нуля")
                    groups=[(f"часть_{j+1}",frame.slice(offset,limit)) for j,offset in enumerate(range(0,frame.height,limit))]
                elif operation=="distribute":
                    count=int(options.get("parts",2))
                    if not 1<=count<=1000:raise ValueError("Количество частей: от 1 до 1000")
                    frame=frame.with_row_index("__source_index")
                    try:col=phone_column(frame,options)
                    except ValueError:col=None
                    if options.get("dedupe",True) and col:frame=frame.unique(subset=[col],maintain_order=True)
                    groups=[(f"часть_{j+1}",frame.filter(pl.col("__source_index")%count==j).drop("__source_index")) for j in range(count)]
                else:
                    if operation=="regions":
                        frame=regions(frame,refs,options); key="Регион РФ"
                        if not options.get("region_zip"):save(frame,name+"_регионы");progress(20+int(i/len(paths)*75));continue
                    else:
                        key=options.get("group_column") or "№ таблицы"
                        if key not in frame.columns:raise ValueError("Нет колонки для группировки: "+key)
                    groups=[(safe(group[key][0]),geo(group,options) if operation=="group_split" else group)
                            for group in frame.partition_by(key,maintain_order=True)]
                parts=[]
                for j,(label,group) in enumerate(groups,1):
                    parts.append(write_table(group,work/(name+"_"+str(j)+"_"+label),options))
                if not parts:parts=[write_table(frame.head(0),work/(name+"_пусто"),options)]
                if operation=="group_split":
                    numbers=set()
                    for _,group in groups:
                        col=phone_column(group,options)
                        for value in group[col].to_list():numbers.update(extract_phones(value))
                    txt=work/(name+"_numbers.txt");txt.write_text("\n".join(sorted(numbers)),encoding="utf-8");parts.append(txt)
                outputs.extend(parts);outputs.append(archive(parts,work/(name+"_части.zip")))
            elif operation=="stats":
                import json
                report=work/(name+"_stats.json");report.write_text(json.dumps(t.basic_stats(frame),ensure_ascii=False,indent=2),encoding="utf-8");outputs.append(report)
            progress(20+int(i/len(paths)*75))
    if len(outputs)>1 and not any(path.suffix==".zip" for path in outputs):
        outputs.append(archive(outputs,work/"все_результаты.zip"))
    return outputs
