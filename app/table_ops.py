from __future__ import annotations
from pathlib import Path
from typing import Iterable
import csv, re, io
from itertools import islice, chain
from html.parser import HTMLParser
import polars as pl

PHONE_ALIASES = {"номер","телефон","phone","mobile","мобильный","тел","phone_number"}
FIO_ALIASES = {"фио","ф.и.о","имя","name","full_name","fullname"}
DATE_ALIASES = {"дата","дата рождения","др","birthdate","birthday","date"}

def _norm_header(x: str) -> str:
    return re.sub(r"\s+", " ", str(x or "").strip().lower().replace("ё","е"))

def _find_col(cols: list[str], aliases: set[str]) -> str | None:
    norm = {_norm_header(c): c for c in cols}
    for a in sorted(aliases, key=len, reverse=True):
        if a in norm:
            return norm[a]
    for c in cols:
        n = _norm_header(c)
        if any(len(a) > 2 and a in n for a in aliases):
            return c
    return None

def _headers(values):
    seen = set()
    result = []
    for i, value in enumerate(values):
        base = str(value).strip() if value is not None else ""
        base = base or f"col_{i+1}"
        name, suffix = base, 2
        while name in seen:
            name = f"{base}_{suffix}"; suffix += 1
        result.append(name); seen.add(name)
    return result

def _from_rows(rows, search_header=False):
    rows = iter(rows)
    if search_header:
        leading=list(islice(rows,20))
        start=next((i for i,row in enumerate(leading)
                    if any(_norm_header(value) in PHONE_ALIASES for value in row)),0)
        rows=chain(leading[start:],rows)
    first = next(rows, [])
    headers = _headers(first)
    if not headers:
        return pl.DataFrame()
    # Treat identifiers as text: leading zeros and large IDs must survive import.
    data = [[None if value is None else str(value) for value in list(row)[:len(headers)]]
            for row in rows if any(value is not None and str(value).strip() for value in row)]
    data = [row + [None] * (len(headers) - len(row)) for row in data]
    return pl.DataFrame(data, schema={name: pl.String for name in headers}, orient="row")

def decode_text(path: Path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Не удалось определить кодировку текста")

class _HTMLTables(HTMLParser):
    def __init__(self):
        super().__init__(); self.tables=[]; self.depth=0; self.rows=[]; self.row=None; self.cell=None
    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.depth += 1
            if self.depth == 1: self.rows=[]
        if self.depth == 1:
            if tag == "tr": self.row=[]
            if tag in ("td", "th"): self.cell=[]
            if tag == "br" and self.cell is not None: self.cell.append(" ")
    def handle_data(self, data):
        if self.cell is not None: self.cell.append(data)
    def handle_endtag(self, tag):
        if self.depth == 1:
            if tag in ("td", "th") and self.cell is not None:
                if self.row is not None: self.row.append("".join(self.cell).strip())
                self.cell=None
            if tag == "tr" and self.row is not None:
                self.rows.append(self.row); self.row=None
            if tag == "table": self.tables.append(self.rows)
        if tag == "table": self.depth=max(0,self.depth-1)

def read_tables(path: Path, first_sheet=False, search_header=False) -> list[tuple[str, pl.DataFrame]]:
    ext = path.suffix.lower()
    if ext == ".csv":
        with path.open("rb") as stream: sample=stream.read(65536)
        try: sample_text=sample.decode("utf-8-sig")
        except UnicodeDecodeError: sample_text=sample.decode("cp1251")
        try: sep=csv.Sniffer().sniff(sample_text, delimiters=",;\t|").delimiter
        except csv.Error: sep=","
        try:
            frame=pl.read_csv(path, separator=sep, infer_schema=False, encoding="utf8",
                              try_parse_dates=False)
        except pl.exceptions.ComputeError as error:
            if "utf-8" not in str(error).lower() and "utf8" not in str(error).lower(): raise
            frame=pl.read_csv(io.StringIO(decode_text(path)), separator=sep, infer_schema=False)
        return [(path.stem, frame)]
    if ext == ".txt":
        return [(path.stem, pl.DataFrame({"Номер": decode_text(path).splitlines()}))]
    if ext in {".xlsx",".xlsm"}:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [(ws.title,_from_rows(ws.iter_rows(values_only=True),search_header)) for ws in
                    (wb.worksheets[:1] if first_sheet else wb.worksheets)]
        finally: wb.close()
    if ext == ".xls":
        import xlrd
        wb=xlrd.open_workbook(path)
        try:
            return [(ws.name,_from_rows((ws.row_values(i) for i in range(ws.nrows)),search_header))
                    for ws in (wb.sheets()[:1] if first_sheet else wb.sheets())]
        finally: wb.release_resources()
    if ext == ".xlsb":
        from pyxlsb import open_workbook
        with open_workbook(str(path)) as wb:
            tables=[]
            for name in (wb.sheets[:1] if first_sheet else wb.sheets):
                with wb.get_sheet(name) as ws:
                    tables.append((name,_from_rows(([cell.v for cell in row] for row in ws.rows()),search_header)))
            return tables
    if ext in {".html",".htm"}:
        parser=_HTMLTables(); parser.feed(decode_text(path))
        if not parser.tables: raise ValueError("В HTML нет таблиц")
        return [(f"Таблица_{i}",_from_rows(rows)) for i,rows in enumerate(parser.tables,1)]
    if ext == ".parquet":
        return [(path.stem,pl.read_parquet(path))]
    raise ValueError(f"Unsupported format: {ext}")

def read_any(path: Path, search_header=False) -> pl.DataFrame:
    tables=read_tables(path,search_header=search_header)
    return pl.concat([frame for _,frame in tables], how="diagonal_relaxed") if tables else pl.DataFrame()

def write_csv(df: pl.DataFrame, out: Path):
    out.parent.mkdir(parents=True, exist_ok=True)
    df.write_csv(out)

def normalize_phone_expr(col: str):
    s = pl.col(col).cast(pl.Utf8, strict=False).fill_null("").str.replace_all(r"\D+","")
    # 10-digit Russian mobile -> prepend 7, 8xxxxxxxxxx -> 7xxxxxxxxxx
    return (
        pl.when(s.str.len_chars()==10).then(pl.lit("7")+s)
        .when((s.str.len_chars()==11) & s.str.starts_with("8")).then(pl.lit("7")+s.str.slice(1))
        .otherwise(s).alias(col)
    )

def normalize_fio_expr(col: str):
    return (
        pl.col(col).cast(pl.Utf8, strict=False).fill_null("")
        .str.replace_all(r"\s+"," ").str.strip_chars()
        .str.to_lowercase().str.to_titlecase().alias(col)
    )

def normalize_dates_expr(col: str):
    s = pl.col(col).cast(pl.Utf8, strict=False).fill_null("").str.strip_chars()
    parsed = pl.coalesce([s.str.strptime(pl.Date, format=fmt, strict=False)
                          for fmt in ("%Y-%m-%d","%d.%m.%Y","%d/%m/%Y","%Y-%m-%d %H:%M:%S")])
    return parsed.dt.strftime("%Y-%m-%d").fill_null(s).alias(col)

def normalize(df: pl.DataFrame) -> pl.DataFrame:
    cols = df.columns
    phone = _find_col(cols, PHONE_ALIASES)
    fio = _find_col(cols, FIO_ALIASES)
    date = _find_col(cols, DATE_ALIASES)
    exprs = []
    if phone: exprs.append(normalize_phone_expr(phone))
    if fio: exprs.append(normalize_fio_expr(fio))
    if date: exprs.append(normalize_dates_expr(date))
    return df.with_columns(exprs) if exprs else df

def clean(df: pl.DataFrame) -> pl.DataFrame:
    """Trim text fields and remove empty text values without changing numeric types."""
    exprs = []
    for name, dtype in df.schema.items():
        if dtype == pl.String:
            value = pl.col(name).str.replace_all(r"\s+", " ").str.strip_chars()
            exprs.append(pl.when(value == "").then(None).otherwise(value).alias(name))
    return df.with_columns(exprs) if exprs else df

def remove_leading_7(df: pl.DataFrame) -> pl.DataFrame:
    phone = _find_col(df.columns, PHONE_ALIASES)
    if not phone:
        raise ValueError("Phone column not found")
    s = pl.col(phone).cast(pl.Utf8, strict=False).fill_null("")
    return df.with_columns(pl.when(s.str.starts_with("7")).then(s.str.slice(1)).otherwise(s).alias(phone))

def dedupe(df: pl.DataFrame, keys: list[str] | None = None) -> pl.DataFrame:
    if keys:
        existing = [k for k in keys if k in df.columns]
    else:
        existing = [c for c in [_find_col(df.columns, PHONE_ALIASES), _find_col(df.columns, {"email","e-mail","почта"}), _find_col(df.columns, {"инн","inn"})] if c]
    return df.unique(subset=existing or None, keep="first", maintain_order=False)

def sort_rows(df: pl.DataFrame, column: str, descending: bool = False) -> pl.DataFrame:
    if not column or column not in df.columns:
        raise ValueError("Укажи колонку для сортировки")
    return df.sort(column, descending=descending, nulls_last=True)

def filter_rows(df: pl.DataFrame, column: str, value: str, contains: bool = False) -> pl.DataFrame:
    if not column or column not in df.columns:
        raise ValueError("Укажи колонку для фильтра")
    if value is None or str(value) == "":
        raise ValueError("Укажи значение для фильтра")
    needle = str(value)
    text = pl.col(column).cast(pl.String, strict=False).fill_null("")
    condition = text.str.contains(needle, literal=True) if contains else text == needle
    return df.filter(condition)

def merge(paths: list[Path]) -> pl.DataFrame:
    frames = []
    all_cols = []
    for p in paths:
        df = read_any(p).with_columns(pl.lit(p.name).alias("_source_file"))
        frames.append(df)
        for c in df.columns:
            if c not in all_cols:
                all_cols.append(c)
    aligned = []
    for df in frames:
        missing = [c for c in all_cols if c not in df.columns]
        if missing:
            df = df.with_columns([pl.lit(None).alias(c) for c in missing])
        aligned.append(df.select(all_cols))
    return pl.concat(aligned, how="vertical_relaxed") if aligned else pl.DataFrame()

def split_csv(df: pl.DataFrame, outdir: Path, max_rows: int = 500_000) -> list[Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    outs = []
    for i, offset in enumerate(range(0, df.height, max_rows), 1):
        p = outdir / f"part_{i:04d}.csv"
        df.slice(offset, max_rows).write_csv(p)
        outs.append(p)
    return outs

def basic_stats(df: pl.DataFrame) -> dict:
    return {"rows": df.height, "columns": df.width, "column_names": df.columns}
