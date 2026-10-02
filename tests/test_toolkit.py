import os
from pathlib import Path
import tempfile
import unittest
import json
import time
import zipfile
import polars as pl
from openpyxl import Workbook, load_workbook
from app import table_ops as t
from app import toolkit_ops as ops


class ToolkitTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.root=Path(self.temp.name)
        self.frame=pl.DataFrame({
            "Номер":["+7 (916) 123-45-67 8 (925) 765-43-21","89991234567","79161234567"],
            "ФИО":["  ИВАНОВ  ИВАН  ИВАНОВИЧ ","Петров Пётр","Сидоров Иван Петрович"],
            "Дата":["2020-01-02","03.02.1990","04/03/1980"],
            "Адрес":["Москва, ул. Ленина, д. 12 к. 2 стр. 1 лит. А кв. 15","Москва","Москва, Ленина, 12"],
            "ID":["001","002","003"],
            "№ таблицы":["1","2","1"],
        })
        self.source=self.root/"data.csv";self.frame.write_csv(self.source,separator=";")
        self.reference=self.root/"max.txt";self.reference.write_text("9161234567;yes\n",encoding="utf8")

    def tearDown(self):
        if hasattr(self,"api_engine"): self.api_engine.dispose()
        self.temp.cleanup()

    def test_csv_sniffing_preserves_identifiers(self):
        imported=t.read_any(self.source)
        self.assertEqual(imported.columns,self.frame.columns)
        self.assertEqual(imported["ID"].to_list(),["001","002","003"])
        encoded=self.root/"cp.csv"
        encoded.write_bytes("Номер;ФИО\n79161234567;Иван Иванов\n".encode("cp1251"))
        self.assertEqual(t.read_any(encoded).width,2)

    def test_phones_and_geo_without_comma(self):
        cases=["+7 (916) 123-45-67 8 (925) 765-43-21","79161234567/79257654321",
               "7916123456779257654321","9161234567 | 9257654321"]
        for text in cases:
            self.assertEqual(ops.extract_phones(text),["79161234567","79257654321"])
        self.assertFalse(ops.extract_phones("Дата 01.01.2020"))
        result=ops.geo(self.frame,{"format_addresses":True})
        self.assertEqual(result.height,4)
        self.assertEqual(result["ФИО"][0],"Иванов Иван Иванович")
        self.assertEqual(result["Дата"][0],"02.01.2020")
        self.assertEqual(result["Адрес требует проверки"].sum(),1)

    def test_max_matches_unmatched_and_optional_seven(self):
        options={"matches_only":True,"download_unmatched":True,"remove7":True,"dedupe":False}
        matching,other=ops.max_filter(self.frame,[self.reference],options)
        self.assertEqual(matching["Номер"].to_list(),["9161234567","9161234567"])
        self.assertEqual(set(other["Номер"].to_list()),{"9257654321","9991234567"})
        self.assertEqual(matching["MAX"].to_list(),["yes","yes"])
        with self.assertRaises(ValueError):ops.max_filter(self.frame,[],options)
        all_rows,_=ops.max_filter(self.frame,[self.reference],{"matches_only":False})
        self.assertEqual(all_rows.height,4)
        century,_=ops.max_filter(self.frame,[self.reference],{"replace_century":True})
        self.assertEqual(century["Дата"][0],"02.01.1920")

    def test_addresses_and_uncertainty(self):
        text,flag=ops.format_address("МОСКВА, улица ЛЕНИНА дом 12 корпус 2 строение 1 литера А квартира 15")
        self.assertEqual(text,"г. Москва, ул. Ленина, д. 12, к. 2, с. 1, лит. А, кв. 15")
        self.assertFalse(flag)
        self.assertEqual(ops.format_address("Москва, ул. Ленина, д. 12к2с1")[0],
                         "г. Москва, ул. Ленина, д. 12, к. 2, с. 1")
        self.assertTrue(ops.format_address("Москва")[1])
        self.assertIn("область",ops.format_address("Московская область, Москва, ул. Ленина, д. 12")[0])

    def test_all_operations_create_readable_results(self):
        region_ref=self.root/"regions.csv"
        region_ref.write_text("Код;От;До;Регион\n916;0;9999999;Москва\n925;0;9999999;Москва\n999;0;9999999;Другой\n",encoding="utf8")
        options={"parts":2,"count":2,"max_rows":2,"column":"ID","value":"001",
                 "matches_only":True,"download_unmatched":True,"region_zip":True}
        for operation in sorted(ops.OPERATIONS):
            with self.subTest(operation=operation):
                outdir=self.root/operation;outdir.mkdir()
                refs=[region_ref] if operation=="regions" else [self.reference] if operation=="max_filter" else []
                paths=[self.source,self.source] if operation=="merge" else [self.source]
                outputs=ops.run_operation(operation,paths,refs,options,outdir,lambda value:None)
                self.assertTrue(outputs)
                for output in outputs:
                    self.assertTrue(output.is_file())
                    if output.suffix==".zip":
                        with zipfile.ZipFile(output) as archive:self.assertIsNone(archive.testzip())
                    if output.suffix==".csv":self.assertTrue(t.read_any(output).width>0)

    def test_excel_and_html(self):
        path=self.root/"book.xlsx";wb=Workbook();ws=wb.active
        ws.append(["Номер","ID"]);ws.append(["79161234567","001"])
        other=wb.create_sheet("Ещё");other.append(["Номер","Город"]);other.append(["79251234567","Москва"]);wb.save(path);wb.close()
        self.assertEqual(t.read_any(path).height,2)
        self.assertEqual(t.read_tables(path,first_sheet=True)[0][1].height,1)
        output=ops.write_table(pl.DataFrame({"text":["=1+1","001"]}),self.root/"result",{"output_format":"xlsx"})
        loaded=load_workbook(output);self.assertEqual(loaded.active["A2"].data_type,"s");loaded.close()
        html=self.root/"table.html";html.write_text("<table><tr><th>ID</th><th>Имя</th></tr><tr><td>001</td><td>Иван</td></tr></table>",encoding="utf8")
        self.assertEqual(t.read_any(html)["ID"][0],"001")

    def test_distribution_and_take_keep_partition(self):
        unique=pl.DataFrame({"Номер":["79161234567","79161234567","79251234567","79991234567"],"ID":["0","1","2","3"]})
        source=self.root/"split.csv";unique.write_csv(source)
        out=self.root/"partition";out.mkdir()
        outputs=ops.run_operation("distribute",[source],[],{"parts":2,"dedupe":True},out,lambda value:None)
        pieces=[t.read_any(path) for path in outputs if path.suffix==".csv"]
        self.assertEqual(sum(frame.height for frame in pieces),3)
        out2=self.root/"take-results";out2.mkdir()
        outputs=ops.run_operation("take",[source],[],{"count":2},out2,lambda value:None)
        frames=[t.read_any(path) for path in outputs if path.suffix==".csv"]
        self.assertEqual([frame.height for frame in frames],[2,2])
        self.assertEqual(set(frames[0]["ID"]) & set(frames[1]["ID"]),set())

    def test_fssp_preamble_sheets_and_formulas(self):
        path=self.root/"fssp.xlsx";wb=Workbook()
        ws=wb.active;ws.append(["Отчёт"]);ws.append(["Номер","ID"]);ws.append(["79161234567","=1+1"])
        other=wb.create_sheet("Второй");other.append(["Номер"]);other.append(["79251234567"]);wb.save(path);wb.close()
        out=self.root/"fssp-out";out.mkdir()
        outputs=ops.run_operation("fssp",[path],[],{"output_format":"xlsx"},out,lambda v:None)
        loaded=load_workbook(outputs[0])
        self.assertEqual(loaded.active["A1"].value,"Отчёт")
        self.assertEqual(loaded.active["A2"].value,"Номер телефона")
        self.assertEqual(loaded.active["B3"].data_type,"f")
        self.assertEqual(loaded["Второй"]["A1"].value,"Номер телефона");loaded.close()
        self.assertEqual(t.read_any(path,search_header=True).height,2)

    def test_api_reference_upload_and_download(self):
        os.environ["DATABASE_URL"]="sqlite:///"+str(self.root/"api.db").replace("\\","/")
        os.environ["WORK_PATH"]=str(self.root/"work")
        os.environ["LOCAL_STORAGE_PATH"]=str(self.root/"storage")
        from app.main import app
        from app.db import engine
        self.api_engine=engine
        from fastapi.testclient import TestClient
        headers={"X-Session-ID":"test-session-001"}
        with TestClient(app) as client:
            response=client.post("/api/quick/jobs",headers=headers,data={
                "operation":"max_filter","options_json":json.dumps({"matches_only":True,"download_unmatched":True,"remove7":True})},
                files=[("files",("data.csv",self.source.read_bytes(),"text/csv")),
                       ("references",("list.txt",self.reference.read_bytes(),"text/plain"))])
            self.assertEqual(response.status_code,200,response.text);jid=response.json()["id"]
            for _ in range(50):
                job=client.get(f"/api/jobs/{jid}",headers=headers).json()
                if job["status"] in ("done","error"):break
                time.sleep(.1)
            self.assertEqual(job["status"],"done",job)
            result=next(item for item in job["outputs"] if item["name"]=="data_MAX.csv")
            download=client.get(f'/api/files/{result["id"]}/download',headers=headers)
            self.assertEqual(download.status_code,200)
            self.assertIn("9161234567",download.text)
            other=client.get(f'/api/files/{result["id"]}/download',headers={"X-Session-ID":"other-session-002"})
            self.assertEqual(other.status_code,404)
            from app.config import settings
            from app.jobs import cleanup_quick_job
            folder=Path(settings.work_path)/f"job_{jid}"
            (folder/result["name"]).unlink()
            self.assertEqual(client.get(f'/api/files/{result["id"]}/download',headers=headers).status_code,404)
            cleanup_quick_job(jid,folder)
            expired=client.get(f"/api/jobs/{jid}",headers=headers).json()
            self.assertTrue(all(output["expired"] for output in expired["outputs"]))
            self.assertFalse(folder.exists())


if __name__=="__main__":unittest.main()
